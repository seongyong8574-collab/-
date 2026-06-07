# -*- coding: utf-8 -*-
"""pptx 도형 좌표/이미지를 파싱해 배치 검증용 PNG 컨택트시트를 생성."""
import io
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

P = Presentation('moet_onview_popup_report_sky.pptx')
SCALE = 96  # px per inch
SW = int(P.slide_width / 914400 * SCALE)
SH = int(P.slide_height / 914400 * SCALE)

def fill_hex(sh):
    try:
        f = sh.fill
        if f.type is not None and f.fore_color and f.fore_color.type is not None:
            return '#' + str(f.fore_color.rgb)
    except Exception:
        pass
    return None

try:
    fnt = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 12)
    fntb = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', 13)
except Exception:
    fnt = fntb = ImageFont.load_default()

slides_img = []
for si, slide in enumerate(P.slides, 1):
    img = Image.new('RGB', (SW, SH), '#ffffff')
    d = ImageDraw.Draw(img, 'RGBA')
    for sh in slide.shapes:
        x = int(sh.left / 914400 * SCALE); y = int(sh.top / 914400 * SCALE)
        w = int(sh.width / 914400 * SCALE); h = int(sh.height / 914400 * SCALE)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                pim = Image.open(io.BytesIO(sh.image.blob)).convert('RGB')
                pim = pim.resize((max(1, w), max(1, h)))
                img.paste(pim, (x, y))
                d.rectangle([x, y, x+w, y+h], outline='#5b9bd5', width=1)
            except Exception as e:
                d.rectangle([x, y, x+w, y+h], outline='red', width=2)
        elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            fh = fill_hex(sh)
            if fh:
                d.rectangle([x, y, x+w, y+h], fill=fh)
        elif sh.has_text_frame:
            # render latin/number text; korean shows as boxes (no font) -> mark region
            t = sh.text_frame.text
            ascii_part = ''.join(ch if ord(ch) < 0x3000 else '·' for ch in t)[:60]
            d.text((x+2, y+1), ascii_part, fill='#444444', font=fnt)
            d.rectangle([x, y, x+max(2,w), y+max(2,h)], outline=(160,160,160,90), width=1)
    d.rectangle([0, 0, SW-1, SH-1], outline='#cccccc', width=1)
    d.text((6, 4), f'Slide {si}', fill='#000000', font=fntb)
    slides_img.append(img)

# 3x3 contact sheet
cols, rows = 3, 3
gap = 16
cw, ch = SW//2, SH//2
sheet = Image.new('RGB', (cols*cw + (cols+1)*gap, rows*ch + (rows+1)*gap), '#e9e9ec')
for i, im in enumerate(slides_img):
    im2 = im.resize((cw, ch))
    r, c = divmod(i, cols)
    sheet.paste(im2, (gap + c*(cw+gap), gap + r*(ch+gap)))
sheet.save('preview_sheet.png')
print('saved preview_sheet.png', sheet.size)
